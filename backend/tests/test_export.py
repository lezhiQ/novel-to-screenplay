import json
import yaml
from backend.app.models import ScreenplayOutput, Scene, Dialogue, Action
from backend.app.export import (
    export_yaml, export_json, export_docx,
    normalize_time, normalize_location, format_scene_heading,
)


def make_sample_screenplay():
    return ScreenplayOutput(
        title="测试剧本",
        scenes=[
            Scene(
                scene_id=1,
                location="INT. 咖啡馆",
                time="日",
                description="开场场景",
                items=[
                    Action(type="action", character="林小雨", action="推开门"),
                    Dialogue(type="dialogue", character="林小雨", line="你好"),
                ],
                dialogues=[
                    Dialogue(character="林小雨", line="你好"),
                ],
                actions=[
                    Action(character="林小雨", action="推开门"),
                ],
            )
        ],
    )


class TestNormalizeTime:
    def test_standard_times(self):
        assert normalize_time("日") == "日"
        assert normalize_time("夜") == "夜"
        assert normalize_time("黄昏") == "黄昏"
        assert normalize_time("清晨") == "清晨"
        assert normalize_time("午后") == "午后"
        assert normalize_time("傍晚") == "傍晚"
        assert normalize_time("深夜") == "深夜"

    def test_non_standard_mapping(self):
        assert normalize_time("白天") == "日"
        assert normalize_time("早晨") == "清晨"
        assert normalize_time("晚上") == "夜"
        assert normalize_time("夜晚") == "夜"
        assert normalize_time("中午") == "午后"
        assert normalize_time("下午") == "午后"

    def test_empty_defaults_to_day(self):
        assert normalize_time("") == "日"
        assert normalize_time(None) == "日"


class TestNormalizeLocation:
    def test_already_has_prefix(self):
        assert normalize_location("INT. 咖啡馆") == "INT. 咖啡馆"
        assert normalize_location("EXT. 操场") == "EXT. 操场"

    def test_lowercase_prefix_normalized(self):
        assert normalize_location("int. 咖啡馆") == "INT. 咖啡馆"
        assert normalize_location("ext. 操场") == "EXT. 操场"

    def test_no_prefix_infers_int(self):
        assert normalize_location("咖啡馆") == "INT. 咖啡馆"
        assert normalize_location("图书馆") == "INT. 图书馆"

    def test_no_prefix_infers_ext(self):
        assert normalize_location("操场") == "EXT. 操场"
        assert normalize_location("公园") == "EXT. 公园"
        assert normalize_location("街上") == "EXT. 街上"

    def test_empty_location(self):
        assert normalize_location("") == "INT. 未知地点"


class TestFormatSceneHeading:
    def test_standard_format(self):
        result = format_scene_heading("INT. 咖啡馆", "日")
        assert result == "INT. 咖啡馆 - 日"

    def test_adds_prefix(self):
        result = format_scene_heading("咖啡馆", "夜")
        assert result == "INT. 咖啡馆 - 夜"

    def test_normalizes_time(self):
        result = format_scene_heading("INT. 咖啡馆", "晚上")
        assert result == "INT. 咖啡馆 - 夜"


class TestExportYaml:
    def test_returns_string(self):
        result = export_yaml(make_sample_screenplay())
        assert isinstance(result, str)

    def test_valid_yaml(self):
        result = export_yaml(make_sample_screenplay())
        data = yaml.safe_load(result)
        assert data["title"] == "测试剧本"
        assert len(data["scenes"]) == 1

    def test_preserves_unicode(self):
        result = export_yaml(make_sample_screenplay())
        assert "咖啡馆" in result
        assert "林小雨" in result

    def test_location_has_prefix(self):
        result = export_yaml(make_sample_screenplay())
        data = yaml.safe_load(result)
        assert data["scenes"][0]["location"].startswith("INT.")

    def test_time_is_normalized(self):
        screenplay = ScreenplayOutput(
            title="测试",
            scenes=[Scene(scene_id=1, location="咖啡馆", time="晚上")],
        )
        result = export_yaml(screenplay)
        data = yaml.safe_load(result)
        assert data["scenes"][0]["time"] == "夜"


class TestExportJson:
    def test_returns_string(self):
        result = export_json(make_sample_screenplay())
        assert isinstance(result, str)

    def test_valid_json(self):
        result = export_json(make_sample_screenplay())
        data = json.loads(result)
        assert data["title"] == "测试剧本"
        assert data["scenes"][0]["scene_id"] == 1

    def test_unicode_not_escaped(self):
        result = export_json(make_sample_screenplay())
        assert "咖啡馆" in result
        assert "\\u" not in result

    def test_location_has_prefix(self):
        result = export_json(make_sample_screenplay())
        data = json.loads(result)
        assert data["scenes"][0]["location"].startswith("INT.")


class TestExportDocx:
    def test_returns_bytes(self):
        result = export_docx(make_sample_screenplay())
        assert isinstance(result, bytes)

    def test_not_empty(self):
        result = export_docx(make_sample_screenplay())
        assert len(result) > 0

    def test_valid_docx(self):
        from docx import Document
        from io import BytesIO
        result = export_docx(make_sample_screenplay())
        doc = Document(BytesIO(result))
        # 文档应有内容
        assert len(doc.paragraphs) > 0

    def test_scene_heading_format(self):
        from docx import Document
        from io import BytesIO
        screenplay = ScreenplayOutput(
            title="测试",
            scenes=[Scene(scene_id=1, location="咖啡馆", time="日")],
        )
        result = export_docx(screenplay)
        doc = Document(BytesIO(result))
        # 检查场景标题格式
        scene_heading_found = False
        for para in doc.paragraphs:
            if "INT." in para.text and "咖啡馆" in para.text and "日" in para.text:
                scene_heading_found = True
                break
        assert scene_heading_found, f"未找到标准场景标题，实际段落: {[p.text for p in doc.paragraphs]}"

    def test_character_name_uppercase(self):
        from docx import Document
        from io import BytesIO
        screenplay = ScreenplayOutput(
            title="测试",
            scenes=[Scene(
                scene_id=1,
                location="INT. 咖啡馆",
                time="日",
                items=[Dialogue(type="dialogue", character="林小雨", line="你好")],
            )],
        )
        result = export_docx(screenplay)
        doc = Document(BytesIO(result))
        # 检查角色名大写
        char_found = False
        for para in doc.paragraphs:
            if "林小雨" in para.text:
                for run in para.runs:
                    if run.bold and run.text == "林小雨":
                        char_found = True
                        break
        assert char_found, "角色名应为大写加粗"
