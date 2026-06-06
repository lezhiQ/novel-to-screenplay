import json
import yaml
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from backend.app.models import ScreenplayOutput


def export_yaml(screenplay: ScreenplayOutput) -> str:
    """导出为 YAML 格式字符串"""
    data = screenplay.model_dump()
    return yaml.dump(data, allow_unicode=True, default_flow_style=False, sort_keys=False)


def export_json(screenplay: ScreenplayOutput) -> str:
    """导出为 JSON 格式字符串"""
    data = screenplay.model_dump()
    return json.dumps(data, ensure_ascii=False, indent=2)


def export_docx(screenplay: ScreenplayOutput) -> bytes:
    """导出为 DOCX 格式字节流"""
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)

    # 标题
    title = doc.add_heading(screenplay.title or "剧本", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 遍历场景
    for scene in screenplay.scenes:
        # 场景标题
        scene_heading = f"场景 {scene.scene_id}: {scene.location or ''}"
        if scene.time:
            scene_heading += f" ({scene.time})"
        doc.add_heading(scene_heading, level=2)

        # 场景描述
        if scene.description:
            p = doc.add_paragraph()
            run = p.add_run(scene.description)
            run.italic = True
            run.font.color.rgb = None  # 默认颜色

        # 处理 items 列表（有序）
        if scene.items:
            for item in scene.items:
                if item.type == "dialogue":
                    p = doc.add_paragraph()
                    # 角色名加粗
                    char_run = p.add_run(f"{item.character}: ")
                    char_run.bold = True
                    # 台词
                    line_run = p.add_run(f'"{item.line}"')
                    # 动作（如果有）
                    if item.action:
                        action_run = p.add_run(f" ({item.action})")
                        action_run.italic = True
                        action_run.font.color.rgb = None
                else:  # action
                    p = doc.add_paragraph()
                    action_text = f"[{item.action}]"
                    if item.character:
                        action_text = f"[{item.character}: {item.action}]"
                    run = p.add_run(action_text)
                    run.italic = True
        else:
            # 兼容旧格式：分别处理 actions 和 dialogues
            for action in scene.actions:
                p = doc.add_paragraph()
                action_text = f"[{action.action}]"
                if action.character:
                    action_text = f"[{action.character}: {action.action}]"
                run = p.add_run(action_text)
                run.italic = True

            for dialogue in scene.dialogues:
                p = doc.add_paragraph()
                char_run = p.add_run(f"{dialogue.character}: ")
                char_run.bold = True
                line_run = p.add_run(f'"{dialogue.line}"')
                if dialogue.action:
                    action_run = p.add_run(f" ({dialogue.action})")
                    action_run.italic = True

        # 场景间分隔
        doc.add_paragraph("─" * 40)

    # 保存到内存
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
